"""Rapports : état du réseau (PDF, Excel) et main courante d'exercice ou d'intervention (PDF)."""
from __future__ import annotations

from datetime import datetime
from xml.sax.saxutils import escape as _x

from .store import Store

STATUTS = {"en-ligne": "En ligne", "degrade": "Dégradé", "hors-ligne": "Hors ligne", "test": "En test", "prevu": "Prévu"}
COLONNES = [("nom", "Nœud"), ("etat", "État"), ("site", "Site"), ("materiel", "Matériel"), ("firmware", "Firmware"),
            ("tx", "TX dBm"), ("bruit", "Bruit dBm"), ("erreursRx", "Erreurs RX"), ("batterie", "Batterie V"),
            ("dispo", "Dispo 30 j"), ("maj", "Dernière lecture")]


def _date(iso) -> str:
    if not iso:
        return ""
    try:
        return datetime.fromisoformat(str(iso)).strftime("%d/%m/%Y %H:%M")
    except ValueError:
        return str(iso)


def _etat(n: dict) -> str:
    return "Voisin" if n.get("proprio") == "externe" else STATUTS.get(n.get("statut"), n.get("statut") or "")


def _lignes(store: Store, voisins: bool) -> list[dict]:
    rank = {"hors-ligne": 0, "degrade": 1, "en-ligne": 2, "test": 3, "prevu": 5}
    nodes = [n for n in store.list_nodes() if voisins or n.get("proprio") != "externe"]
    nodes.sort(key=lambda n: (n.get("proprio") == "externe", rank.get(n.get("statut"), 4), n.get("nom", "")))
    out = []
    for n in nodes:
        out.append({
            "nom": n.get("nom", ""), "etat": _etat(n),
            "site": " · ".join(x for x in (n.get("site"), n.get("commune")) if x),
            "materiel": n.get("materiel", ""), "firmware": n.get("firmware", ""),
            "tx": n.get("tx", ""), "bruit": n.get("bruit", ""), "erreursRx": n.get("erreursRx", ""),
            "batterie": n.get("batterie", ""), "maj": _date(n.get("maj")),
            "dispo": (f"{store.availability(n['id'], 30)['pct']} %" if store.availability(n['id'], 30)['pct'] is not None else "")
                     if n.get("proprio") != "externe" and n.get("statut") != "prevu" else "",
            "aFaire": n.get("aFaire", ""),
        })
    return out


def _synthese(store: Store) -> dict:
    own = [n for n in store.list_nodes() if n.get("proprio") != "externe"]
    s = {v: 0 for v in STATUTS.values()}
    for n in own:
        s[STATUTS.get(n.get("statut"), "Prévu")] = s.get(STATUTS.get(n.get("statut"), "Prévu"), 0) + 1
    return s


# ---------- Excel ----------
def excel_etat(store: Store, chemin: str, voisins: bool = False):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Nœuds"
    ws["A1"] = "ADRASEC 06 — État du réseau MeshCore"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = f"Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')}"
    head = [t for _, t in COLONNES] + ["À faire"]
    ws.append([])
    ws.append(head)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D3A57")
        c.alignment = Alignment(vertical="center")
    fills = {"Hors ligne": "F8D7D7", "Dégradé": "FCEBC7", "En ligne": "DDF1E3"}
    for row in _lignes(store, voisins):
        ws.append([row[k] for k, _ in COLONNES] + [row["aFaire"]])
        f = fills.get(row["etat"])
        if f:
            ws.cell(ws.max_row, 2).fill = PatternFill("solid", fgColor=f)
    widths = [30, 12, 34, 16, 11, 9, 11, 11, 11, 11, 18, 50]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A5"

    js = wb.create_sheet("Journal")
    js.append(["Date", "Nœud", "Opérateur", "Texte"])
    for c in js[1]:
        c.font = Font(bold=True)
    for j in store.list_journal(5000):
        js.append([_date(j["date"]), j.get("noeud", ""), j.get("auteur", ""), j.get("texte", "")])
    for col, w in zip("ABCD", (18, 30, 12, 100)):
        js.column_dimensions[col].width = w
    wb.save(chemin)


# ---------- PDF ----------
def _pdf_base(chemin: str, titre: str):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate
    from reportlab.lib.units import mm

    def entete(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.setFillGray(0.4)
        canvas.drawString(15 * mm, 8 * mm, f"ADRASEC 06 — {titre}")
        canvas.drawRightString(doc.pagesize[0] - 15 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(chemin, pagesize=landscape(A4), leftMargin=15 * mm, rightMargin=15 * mm,
                            topMargin=14 * mm, bottomMargin=16 * mm, title=titre, author="ADRASEC 06")
    return doc, entete


def _styles():
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    ss = getSampleStyleSheet()
    return {
        "titre": ParagraphStyle("t", parent=ss["Title"], fontSize=18, alignment=0, spaceAfter=2),
        "sous": ParagraphStyle("s", parent=ss["Normal"], fontSize=9, textColor="#555555", spaceAfter=10),
        "h2": ParagraphStyle("h", parent=ss["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=6),
        "cell": ParagraphStyle("c", parent=ss["Normal"], fontSize=8, leading=10),
        "norm": ParagraphStyle("n", parent=ss["Normal"], fontSize=9, leading=12),
    }


def _table(data, widths, couleurs_etat=None):
    from reportlab.lib import colors
    from reportlab.platypus import Table, TableStyle
    t = Table(data, colWidths=widths, repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1D3A57")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#C9D1D9")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3F6F8")]),
    ]
    for (r, c, col) in (couleurs_etat or []):
        style.append(("BACKGROUND", (c, r), (c, r), colors.HexColor(col)))
    t.setStyle(TableStyle(style))
    return t


def pdf_etat(store: Store, chemin: str, voisins: bool = False):
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, Spacer
    st = _styles()
    doc, entete = _pdf_base(chemin, "État du réseau MeshCore")
    elems = [Paragraph("ADRASEC 06 — État du réseau MeshCore", st["titre"]),
             Paragraph(f"Édité le {datetime.now().strftime('%d/%m/%Y à %H:%M')} · Fréquence 869,618 MHz · "
                       f"BW 62,5 kHz · SF 8 · CR 4/8", st["sous"])]
    syn = _synthese(store)
    elems.append(Paragraph("Synthèse des nœuds ADRASEC : " + " · ".join(f"<b>{v}</b> {k.lower()}" for k, v in syn.items()),
                           st["norm"]))
    elems.append(Spacer(1, 6))
    head = [t for _, t in COLONNES]
    data = [head]
    marks = []
    cols = {"Hors ligne": "#F8D7D7", "Dégradé": "#FCEBC7", "En ligne": "#DDF1E3"}
    lignes = _lignes(store, voisins)
    for i, r in enumerate(lignes, 1):
        data.append([Paragraph(_x(str(r[k] if r[k] is not None else "")), st["cell"]) for k, _ in COLONNES])
        if r["etat"] in cols:
            marks.append((i, 1, cols[r["etat"]]))
    w = [48, 19, 46, 24, 17, 13, 16, 17, 16, 17, 27]
    elems.append(_table(data, [x * mm for x in w], marks))

    todo = [r for r in lignes if r["aFaire"]]
    if todo:
        elems.append(Paragraph("Actions en attente", st["h2"]))
        data = [["Nœud", "À faire"]] + [[Paragraph(_x(r["nom"]), st["cell"]), Paragraph(_x(r["aFaire"]), st["cell"])] for r in todo]
        elems.append(_table(data, [60 * mm, 207 * mm]))

    alertes = store.list_alerts(20)
    if alertes:
        elems.append(Paragraph("Dernières alertes", st["h2"]))
        data = [["Date", "Nœud", "Alerte"]] + [[_date(a["date"]), Paragraph(_x(a.get("noeud") or ""), st["cell"]),
                                                Paragraph(_x(a["texte"]), st["cell"])] for a in alertes]
        elems.append(_table(data, [32 * mm, 60 * mm, 175 * mm]))
    doc.build(elems, onFirstPage=entete, onLaterPages=entete)


def pdf_main_courante(store: Store, ex_id: int, chemin: str):
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph
    st = _styles()
    ex = store.get_exercise(ex_id)
    if not ex:
        raise ValueError("Exercice introuvable.")
    titre = f"Main courante — {ex['nom']}"
    doc, entete = _pdf_base(chemin, titre)
    elems = [Paragraph(_x(f"ADRASEC 06 — {titre}"), st["titre"]),
             Paragraph(f"{(ex.get('type') or 'exercice').capitalize()} · Début {_date(ex['debut'])} · "
                       f"Fin {_date(ex['fin']) if ex.get('fin') else 'en cours'} · Responsable {ex.get('responsable') or '—'}",
                       st["sous"])]
    lignes = store.journal_of_exercise(ex_id)
    data = [["Heure", "Opérateur", "Nœud", "Événement"]]
    for j in lignes:
        data.append([_date(j["date"]), j.get("auteur") or "", Paragraph(_x(j.get("noeud") or ""), st["cell"]),
                     Paragraph(_x(j["texte"]), st["cell"])])
    elems.append(_table(data, [30 * mm, 22 * mm, 55 * mm, 160 * mm]))

    t0, t1 = ex["debut"], ex.get("fin") or "9999"
    msgs = [m for m in store.list_messages(5000) if t0 <= m["date"] <= t1]
    if msgs:
        elems.append(Paragraph("Messages radio échangés", st["h2"]))
        data = [["Heure", "Sens", "Canal / contact", "Message"]]
        for m in msgs:
            ou = m.get("canal_nom") or m.get("contact") or ""
            data.append([_date(m["date"]), "Reçu" if m["sens"] == "recu" else "Émis", Paragraph(_x(ou), st["cell"]),
                         Paragraph(_x((m.get("auteur") + " : " if m.get("auteur") and m["sens"] == "recu" else "") + m["texte"]), st["cell"])])
        elems.append(_table(data, [30 * mm, 14 * mm, 50 * mm, 173 * mm]))
    doc.build(elems, onFirstPage=entete, onLaterPages=entete)
